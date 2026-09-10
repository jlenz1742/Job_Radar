# -*- coding: utf-8 -*-
"""
Job-Radar — Zustands- und Auswertungsmaschine.

Das Skript ruft NICHTS ab. Das Abrufen (Karriereseiten, jobs.ch, Alert-Mails)
macht der Agent über WebFetch bzw. den Gmail-Konnektor und übergibt die Funde
als JSON. Dieses Skript macht den deterministischen Teil:
Normalisieren -> Deduplizieren -> Zustand fortschreiben -> Excel bauen.

Aufruf:
    python radar.py ingest funde.json        # Funde eines Laufs einlesen
    python radar.py excel Job_Radar.xlsx     # Arbeitsmappe bauen
    python radar.py neu                      # Neuzugänge des letzten Laufs
"""

import json, os, re, sys, unicodedata
from datetime import datetime, timezone

HERE = os.path.dirname(os.path.abspath(__file__))
STATE = os.path.join(HERE, "state.json")

# ---------------------------------------------------------------- Radius

DRIN = {"ZH","ZG","SZ","AG","SH","TG","SG","AR","AI","LU","NW","OW","UR",
        "GL","BL","BS","SO","BE","FL"}
DRAUSSEN = {"GE","VD","VS","NE","JU","FR","TI","GR"}

def radius(kanton):
    k = (kanton or "").strip().upper()
    if k in DRIN:
        return "drin"
    if k in DRAUSSEN:
        return "weiter weg"
    return "unbekannt"

# ---------------------------------------------------------------- Schlüssel

# Alles, was denselben Job unter anderem Namen erscheinen lässt.
RAUSCH = [
    r"\(\s*[amwdfxh/\s\*:]+\s*\)",      # (m/w/d), (a), (f/m/d), (w/m/d)
    r"\b\d{1,3}\s*[-–]\s*\d{1,3}\s*%",  # 80-100%
    r"\b\d{1,3}\s*%",                   # 100%
    r"\bm/w/d\b|\bf/m/d\b|\bw/m/d\b|\bm/f/d\b|\bh/f/x\b|\bh/f\b",
    r"\bvollzeit\b|\bteilzeit\b|\bfulltime\b|\bfull time\b|\bpart time\b",
    r"\bbefristet\b|\bunbefristet\b|\bpermanent\b|\btemporary\b",
]

RECHTSFORM = [r"\bag\b", r"\bsa\b", r"\bgmbh\b", r"\bsarl\b", r"\bsàrl\b",
              r"\bholding\b", r"\bgroup\b", r"\bgruppe\b", r"\bschweiz\b",
              r"\bswitzerland\b", r"\bsuisse\b", r"\binternational\b",
              r"\bs\.?a\.?u\.?\b", r"\b\(schweiz\)\b"]

def entumlauten(s):
    s = s.replace("ä","ae").replace("ö","oe").replace("ü","ue").replace("ß","ss")
    s = s.replace("Ä","Ae").replace("Ö","Oe").replace("Ü","Ue")
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()

def norm(s, rechtsform_weg=False):
    s = entumlauten((s or "").lower())
    for p in RAUSCH:
        s = re.sub(p, " ", s)
    if rechtsform_weg:
        for p in RECHTSFORM:
            s = re.sub(p, " ", s)
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()

def schluessel(firma, titel):
    """Firma + Titel, unscharf. Fängt dieselbe Stelle aus drei Quellen ein."""
    return norm(firma, True) + "|" + norm(titel)

# ---------------------------------------------------------------- Zustand

def leer():
    return {"version": 1, "laeufe": [], "jobs": {}, "quellen": {}}

def laden():
    if not os.path.exists(STATE):
        return leer()
    with open(STATE, encoding="utf-8") as f:
        return json.load(f)

def sichern(st):
    with open(STATE, "w", encoding="utf-8") as f:
        json.dump(st, f, ensure_ascii=False, indent=1)

FELDER = ["firma","titel","ort","kanton","quelle","quelle_id","url","publiziert",
          "prio","teil","bewertung","consulting","notiz"]

def ingest(pfad):
    with open(pfad, encoding="utf-8") as f:
        daten = json.load(f)
    funde  = daten.get("funde", [])
    quellen = daten.get("quellen", {})     # {"Tecan": 38, "Hilti": "FEHLER"}
    heute  = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    st = laden()
    erstlauf = not st["laeufe"]

    gesehen, neu, wieder = set(), [], []
    for r in funde:
        k = schluessel(r.get("firma",""), r.get("titel",""))
        if not k.strip("|"):
            continue
        gesehen.add(k)
        j = st["jobs"].get(k)
        if j is None:
            j = {f: r.get(f, "") for f in FELDER}
            j["quelle_id"] = r.get("quelle_id") or r.get("firma", "")
            j.update(key=k, erstmals=heute, zuletzt=heute,
                     status="neu", laeufe_gesehen=1)
            # Beim allerersten Lauf ist nichts "neu" — es ist der Bestand.
            if erstlauf:
                j["status"] = "offen"
            st["jobs"][k] = j
            (neu if not erstlauf else wieder).append(j)
        else:
            # Bestehende Zeile: nur ergänzen, nie überschreiben, was schon dasteht.
            for f in FELDER:
                if r.get(f) and not j.get(f):
                    j[f] = r[f]
            j["zuletzt"] = heute
            j["laeufe_gesehen"] = j.get("laeufe_gesehen", 1) + 1
            if j["status"] in ("neu", "weg"):
                j["status"] = "offen"

    # "Weg" darf eine Stelle nur werden, wenn ihre EIGENE Quelle in diesem Lauf
    # abgefragt wurde UND sauber geliefert hat. Ein Lauf, der nur fünf Quellen
    # abdeckt, darf nicht den ganzen Bestand für tot erklären.
    verdaechtig = {q for q, v in quellen.items()
                   if isinstance(v, int) and v == 0
                   and isinstance(st["quellen"].get(q, {}).get("anzahl"), int)
                   and st["quellen"][q]["anzahl"] >= 5}
    gesunde = {q for q, v in quellen.items()
               if isinstance(v, int) and q not in verdaechtig}
    verschwunden = []
    for k, j in st["jobs"].items():
        if k in gesehen or j["status"] == "weg":
            continue
        qid = j.get("quelle_id") or j.get("firma", "")
        if qid not in gesunde:
            continue          # Quelle nicht abgefragt oder kaputt — kein Urteil
        j["status"] = "weg"
        j["weg_seit"] = heute
        verschwunden.append(j)

    # Quellen-Wächter: stiller Ausfall ist der gefährlichste Fehler.
    alarme = []
    for q, v in quellen.items():
        alt = st["quellen"].get(q, {})
        vorher = alt.get("anzahl")
        if isinstance(v, int):
            if isinstance(vorher, int) and vorher >= 5 and v == 0:
                alarme.append(f"{q}: 0 Stellen, letztes Mal {vorher}")
            st["quellen"][q] = {"anzahl": v, "ok": heute,
                                "letzter_fehler": alt.get("letzter_fehler","")}
        else:
            alarme.append(f"{q}: {v}")
            st["quellen"][q] = {"anzahl": vorher, "ok": alt.get("ok",""),
                                "letzter_fehler": heute}

    st["laeufe"].append({
        "ts": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "funde": len(funde), "neu": len(neu), "weg": len(verschwunden),
        "bestand": sum(1 for j in st["jobs"].values() if j["status"] != "weg"),
        "alarme": alarme,
    })
    sichern(st)

    print(f"Lauf {len(st['laeufe'])} | eingelesen {len(funde)} | "
          f"neu {len(neu)} | weg {len(verschwunden)} | "
          f"Bestand {st['laeufe'][-1]['bestand']}")
    if erstlauf:
        print(f"Erstlauf: {len(wieder)} Stellen als Bestand aufgenommen (nichts als 'neu' markiert).")
    for j in neu:
        print(f"  NEU  {j['firma']} — {j['titel']} ({j['ort']})")
    for j in verschwunden:
        print(f"  WEG  {j['firma']} — {j['titel']}")
    for a in alarme:
        print(f"  !!   {a}")

def neu_zeigen():
    st = laden()
    if not st["laeufe"]:
        print("Noch kein Lauf.")
        return
    letzter = st["laeufe"][-1]["ts"][:10]
    n = [j for j in st["jobs"].values()
         if j.get("status") == "neu" and j["erstmals"] == letzter]
    print(f"Neu seit dem Lauf davor ({letzter}): {len(n)}")
    for j in sorted(n, key=lambda x: (x.get("prio",""), x["firma"])):
        print(f"  [{j.get('prio','—')}] {j['firma']} — {j['titel']} "
              f"({j['ort']}, {radius(j.get('kanton'))}) {j['url']}")

# ---------------------------------------------------------------- Excel

def excel(pfad):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.comments import Comment

    ARIAL = "Arial"
    KOPF  = PatternFill("solid", fgColor="1F3864")
    F_NEU = PatternFill("solid", fgColor="FFF2CC")
    F_WEG = PatternFill("solid", fgColor="F2F2F2")
    F_EIN = PatternFill("solid", fgColor="FFFF00")
    THIN  = Side(style="thin", color="D0D0D0")
    RAND  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    st = laden()
    letzter = st["laeufe"][-1]["ts"][:10] if st["laeufe"] else ""

    SPALTEN = [
        ("Firma", 26), ("Stellentitel", 46), ("Ort", 20), ("KT", 5),
        ("Radius", 12), ("Prio", 6), ("Teil", 6), ("Quelle", 16),
        ("Publiziert", 12), ("Erstmals gesehen", 15), ("Zuletzt gesehen", 15),
        ("Läufe", 7), ("Status", 9), ("Bewertung", 11), ("Consulting", 11),
        ("Mein Status", 14), ("Meine Notiz", 34), ("Link", 46),
    ]
    EINGABE = {"Mein Status", "Meine Notiz"}

    def kopfzeile(ws):
        for i, (name, breite) in enumerate(SPALTEN, 1):
            c = ws.cell(1, i, name)
            c.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
            c.fill = KOPF
            c.alignment = Alignment(vertical="center", horizontal="center",
                                    wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = breite
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "A2"

    def zeile(ws, r, j):
        werte = [
            j.get("firma",""), j.get("titel",""), j.get("ort",""),
            (j.get("kanton","") or "").upper(), radius(j.get("kanton")),
            j.get("prio","") or "—", j.get("teil","") or "",
            j.get("quelle",""), j.get("publiziert",""),
            j.get("erstmals",""), j.get("zuletzt",""),
            j.get("laeufe_gesehen",1), j.get("status",""),
            j.get("bewertung",""), j.get("consulting",""),
            j.get("mein_status",""), j.get("notiz",""), j.get("url",""),
        ]
        for i, v in enumerate(werte, 1):
            c = ws.cell(r, i, v)
            c.font = Font(name=ARIAL, size=9)
            c.border = RAND
            c.alignment = Alignment(vertical="top", wrap_text=(i in (2, 17)))
            if SPALTEN[i-1][0] in EINGABE:
                c.fill = F_EIN
            elif j.get("status") == "neu":
                c.fill = F_NEU
            elif j.get("status") == "weg":
                c.fill = F_WEG
        if j.get("url"):
            lz = ws.cell(r, len(SPALTEN))
            lz.hyperlink = j["url"]
            lz.font = Font(name=ARIAL, size=9, color="0563C1", underline="single")

    wb = openpyxl.Workbook()

    # --- Blatt 1: Neuzugänge des letzten Laufs -------------------------
    ws = wb.active
    ws.title = "Neu"
    kopfzeile(ws)
    # Nur echte Neuzugänge: Status "neu" wird beim Erstlauf bewusst nicht gesetzt.
    neu = [j for j in st["jobs"].values()
           if j.get("status") == "neu" and j.get("erstmals") == letzter]
    neu.sort(key=lambda x: (str(x.get("prio","9")), x.get("firma","")))
    for r, j in enumerate(neu, 2):
        zeile(ws, r, j)
    ws.cell(1, 1).comment = Comment(
        f"Neuzugänge des Laufs vom {letzter}: {len(neu)} Stellen.\n"
        "Beim Erstlauf ist dieses Blatt leer — dort ist alles Bestand.\n"
        "Gelbe Spalten sind für dich: Mein Status, Meine Notiz.",
        "Job-Radar", height=120, width=320)

    # --- Blatt 2: alle Stellen ----------------------------------------
    ws = wb.create_sheet("Alle Stellen")
    kopfzeile(ws)
    alle = sorted(st["jobs"].values(),
                  key=lambda x: (x.get("status") == "weg",
                                 x.get("erstmals",""), x.get("firma","")),
                  reverse=False)
    alle.sort(key=lambda x: (x.get("status") == "weg", x.get("erstmals","")),
              reverse=True)
    for r, j in enumerate(alle, 2):
        zeile(ws, r, j)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SPALTEN))}{max(2, len(alle)+1)}"

    # --- Blatt 3: Quellen-Wächter -------------------------------------
    ws = wb.create_sheet("Quellen")
    for i, (name, breite) in enumerate(
            [("Quelle", 30), ("Stellen letzter OK-Lauf", 22),
             ("Zuletzt OK", 14), ("Letzter Fehler", 14), ("Zustand", 34)], 1):
        c = ws.cell(1, i, name)
        c.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        c.fill = KOPF
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.freeze_panes = "A2"
    for r, (q, v) in enumerate(sorted(st["quellen"].items()), 2):
        zustand = "ok"
        if v.get("letzter_fehler") and v["letzter_fehler"] >= (v.get("ok") or ""):
            zustand = "FEHLER — Quelle liefert nicht"
        elif v.get("anzahl") == 0:
            zustand = "0 Stellen — prüfen"
        for i, val in enumerate([q, v.get("anzahl"), v.get("ok",""),
                                 v.get("letzter_fehler",""), zustand], 1):
            c = ws.cell(r, i, val)
            c.font = Font(name=ARIAL, size=9,
                          bold=zustand.startswith("FEHLER"),
                          color="C00000" if zustand.startswith("FEHLER") else "000000")
            c.border = RAND
    ws.cell(1, 1).comment = Comment(
        "Der wichtigste Blick vor dem Lesen der Trefferliste.\n"
        "Eine Quelle, die antwortet aber nichts liefert, sieht aus wie "
        "'keine neuen Stellen' — und ist der gefährlichste Fehler des Radars.",
        "Job-Radar", height=110, width=320)

    # --- Blatt 4: Laufhistorie ----------------------------------------
    ws = wb.create_sheet("Läufe")
    for i, (name, breite) in enumerate(
            [("Lauf", 6), ("Zeitpunkt (UTC)", 22), ("Eingelesen", 12),
             ("Neu", 8), ("Weg", 8), ("Bestand", 10), ("Alarme", 60)], 1):
        c = ws.cell(1, i, name)
        c.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        c.fill = KOPF
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.freeze_panes = "A2"
    for r, l in enumerate(st["laeufe"], 2):
        for i, val in enumerate([r-1, l["ts"], l["funde"], l["neu"], l["weg"],
                                 l["bestand"], " · ".join(l.get("alarme", []))], 1):
            c = ws.cell(r, i, val)
            c.font = Font(name=ARIAL, size=9)
            c.border = RAND
            c.alignment = Alignment(vertical="top", wrap_text=(i == 7))

    wb.save(pfad)
    print(f"{pfad} geschrieben | {len(alle)} Stellen "
          f"({sum(1 for j in alle if j.get('status') != 'weg')} offen) | "
          f"{len(neu)} neu | {len(st['quellen'])} Quellen")

# ---------------------------------------------------------------- CLI

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(1)
    cmd = sys.argv[1]
    if cmd == "ingest":
        ingest(sys.argv[2])
    elif cmd == "excel":
        excel(sys.argv[2] if len(sys.argv) > 2 else "Job_Radar.xlsx")
    elif cmd == "neu":
        neu_zeigen()
    else:
        print(__doc__); sys.exit(1)
